#!/usr/bin/env python3
"""
build_data.py — 將 Quectel 天線 Excel 選型表轉為 PWA 內嵌資料，並產生 index.html

用法:
    python3 build_data.py \
        --version     2.0.0 \
        --global-xlsx Quectel_Antenna_Product_Selector_V3_7_20260916.xlsx \
        --jp-xlsx     Quectel_Antenna_Product_for_JP_202604__-_update.xlsx \
        --template    template.html \
        --out         index.html

輸出:
    index.html    內嵌資料、版本號與建置時間的 App
    version.json  供 App 檢查伺服器上是否有新版本
    sw.js         自動將 CACHE_VERSION 改為 antenna-selector-v<版本號>，
                  不需再手動修改

資料來源對應:
    Global 頁  : Selector 檔的 "PRO" + "Non-PRO" 工作表
    日本市場頁 : JP 檔的 "5G" / "LTE (4G)" / "Wi-Fi & BT" 工作表
                 + Selector 檔 "PRO-JP" 工作表中有勾選模組相容(√)的料號
                 模組相容性取兩份來源的聯集。

刻意排除的欄位 (內部資訊，不放入可分享的工具):
    - "Buy&Sell/Self-Production(only PDMs know)"
    - Datasheet 內部 SharePoint / OneDrive 連結 (前端可另設 DATASHEET_URL 樣板)

相依套件: openpyxl, Pillow (含 WebP 支援)
"""
from __future__ import annotations

import argparse
import datetime as dt
import base64
import io
import json
import re
import sys
import zipfile
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

THUMB_PX = 160          # 縮圖最長邊 (px)
THUMB_QUALITY = 72      # WebP 品質

EMPTY = {None, "", "None", "N/A", "n/a", "NA", "-"}


# --------------------------------------------------------------------------- #
# 通用工具
# --------------------------------------------------------------------------- #
def clean(v):
    """將儲存格值正規化為字串；空值回傳 None。"""
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).replace("\\n", "\n").replace("\u2011", "-").strip()
    s = re.sub(r"[ \t]+", " ", s)
    return None if s in EMPTY else s


def norm_module(name: str) -> str:
    """模組名稱正規化：不換行連字號→一般連字號、去除換行與多餘空白。"""
    s = str(name).replace("\u2011", "-").replace("\n", " ")
    return re.sub(r"\s+", " ", s).strip()


_RANGE_RE = re.compile(
    r"(DC|\d+(?:\.\d+)?)\s*(GHz|MHz)?\s*(?:[–—\-~]|to)\s*(\d+(?:\.\d+)?)\s*(GHz|MHz)?",
    re.I,
)
_SINGLE_RE = re.compile(r"(\d{3,5}(?:\.\d+)?)\s*(GHz|MHz)?", re.I)


def parse_ranges(text: str | None) -> list[list[float]]:
    """
    從頻率描述文字解析出 [low, high] (MHz) 區間清單並合併重疊。
    支援 '617–960 MHz'、'5.15-5.85GHz'、'DC~6GHz'、'1575.42 MHz' 等寫法。
    假設：未標單位的數值視為 MHz；數值 < 10 且未標單位則視為 GHz。
    """
    if not text:
        return []
    out: list[list[float]] = []
    consumed = []
    for m in _RANGE_RE.finditer(text):
        lo_s, u1, hi_s, u2 = m.groups()
        unit = (u2 or u1 or "").lower()
        lo = 0.0 if lo_s.upper() == "DC" else float(lo_s)
        hi = float(hi_s)
        mul_hi = 1000 if unit == "ghz" or (not unit and hi < 10) else 1
        mul_lo = 1000 if (u1 or "").lower() == "ghz" or (not u1 and unit == "ghz") or (not unit and lo < 10 and lo > 0) else 1
        lo, hi = lo * mul_lo, hi * mul_hi
        if hi > lo:
            out.append([round(lo, 2), round(hi, 2)])
            consumed.append(m.span())
    # 單點頻率 (例如 GNSS 1575.42 MHz) 以 ±1 MHz 表示
    masked = list(text)
    for a, b in consumed:
        for i in range(a, b):
            masked[i] = " "
    for m in _SINGLE_RE.finditer("".join(masked)):
        f = float(m.group(1)) * (1000 if (m.group(2) or "").lower() == "ghz" else 1)
        if 100 <= f <= 100000:
            out.append([f - 1, f + 1])
    out.sort()
    merged: list[list[float]] = []
    for lo, hi in out:
        if merged and lo <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], hi)
        else:
            merged.append([lo, hi])
    return merged


# --------------------------------------------------------------------------- #
# 圖片擷取：解析 drawing XML，以錨點列號對應圖片
# --------------------------------------------------------------------------- #
class SheetImages:
    def __init__(self, xlsx_path: Path):
        self.z = zipfile.ZipFile(xlsx_path)
        wb_xml = self.z.read("xl/workbook.xml").decode("utf8")
        rels = self._rels("xl/_rels/workbook.xml.rels")
        self.sheet_file = {}
        for name, rid in re.findall(r'<sheet [^>]*?name="([^"]+)"[^>]*?r:id="([^"]+)"', wb_xml):
            name = name.replace("&amp;", "&")
            self.sheet_file[name] = "xl/" + rels[rid].lstrip("/").replace("xl/", "")

    def _rels(self, path):
        try:
            xml = self.z.read(path).decode("utf8")
        except KeyError:
            return {}
        res = {}
        for tag in re.findall(r"<Relationship [^>]+>", xml):
            rid = re.search(r'Id="([^"]+)"', tag).group(1)
            tgt = re.search(r'Target="([^"]+)"', tag).group(1)
            res[rid] = tgt
        return res

    def row_images(self, sheet: str) -> dict[int, bytes]:
        """回傳 {0-based 列號: 圖片 bytes}，同列多張僅取第一張。"""
        sf = self.sheet_file.get(sheet)
        if not sf:
            return {}
        base = sf.rsplit("/", 1)
        srels = self._rels(f"{base[0]}/_rels/{base[1]}.rels")
        drawing = next((t for t in srels.values() if "drawings/" in t), None)
        if not drawing:
            return {}
        dpath = "xl/drawings/" + drawing.split("drawings/")[-1]
        dname = dpath.rsplit("/", 1)
        drels = self._rels(f"{dname[0]}/_rels/{dname[1]}.rels")
        xml = self.z.read(dpath).decode("utf8")
        result: dict[int, bytes] = {}
        for anchor in re.findall(r"<xdr:(?:twoCellAnchor|oneCellAnchor)\b.*?</xdr:(?:twoCellAnchor|oneCellAnchor)>", xml, re.S):
            row = re.search(r"<xdr:from>.*?<xdr:row>(\d+)</xdr:row>", anchor, re.S)
            emb = re.search(r'r:embed="([^"]+)"', anchor)
            if not row or not emb or emb.group(1) not in drels:
                continue
            r = int(row.group(1))
            if r in result:
                continue
            media = "xl/media/" + drels[emb.group(1)].split("media/")[-1]
            try:
                result[r] = self.z.read(media)
            except KeyError:
                pass
        return result


def to_thumb(raw: bytes) -> str | None:
    """轉為 WebP 縮圖 data URI；失敗回傳 None。"""
    try:
        im = Image.open(io.BytesIO(raw))
        im = im.convert("RGBA")
        # 裁掉白邊/透明邊，讓產品主體更大
        bg = Image.new("RGBA", im.size, (255, 255, 255, 255))
        flat = Image.alpha_composite(bg, im).convert("RGB")
        inv = Image.eval(flat, lambda p: 255 - p)
        bbox = inv.point(lambda p: 255 if p > 12 else 0).getbbox()
        if bbox:
            flat = flat.crop(bbox)
        flat.thumbnail((THUMB_PX, THUMB_PX), Image.LANCZOS)
        buf = io.BytesIO()
        flat.save(buf, "WEBP", quality=THUMB_QUALITY, method=6)
        return "data:image/webp;base64," + base64.b64encode(buf.getvalue()).decode()
    except Exception:
        return None


# --------------------------------------------------------------------------- #
# Global 選型表
# --------------------------------------------------------------------------- #
def read_selector_sheet(ws, imgs: dict[int, bytes], source: str, images: dict):
    rows = list(ws.iter_rows(values_only=True))
    hdr = [clean(h) for h in rows[0]]
    idx = {h: i for i, h in enumerate(hdr) if h}
    oc_col = idx.get("OC", 0)
    section = None
    products = []
    for rnum, r in enumerate(rows[2:], start=2):
        get = lambda k: clean(r[idx[k]]) if k in idx and idx[k] < len(r) else None
        oc = clean(r[oc_col])
        name = get("Product Name")
        if oc and not name:
            section = oc          # Non-PRO 的分段標題列 (5G / 4G / WIFI ...)
            continue
        if not oc or not name:
            continue
        techs = []
        for n in range(1, 5):
            t = get(f"Technology{n}")
            if not t:
                continue
            techs.append({
                "t": t, "f": get(f"Frequency Range{n}"), "q": get(f"Quantity{n}"),
                "eff": get(f"Efficiency{n}"), "gain": get(f"Peak Gain{n}"),
                "pat": get(f"Radiation Pattern{n}"), "pol": get(f"Polarization{n}"),
            })
        freq_text = get("Frequency Range") or "\n".join(t["f"] or "" for t in techs)
        cables = [c for c in (
            {"len": get("Cable Length1"), "type": get("Cable Type1")},
            {"len": get("Cable Length2"), "type": get("Cable Type2")},
        ) if c["len"] or c["type"]]
        p = {
            "oc": oc, "src": source,
            "type": get("Type") or section,
            "brochure": get("Brochure Type"),
            "name": name, "desc": get("Detailed Description"),
            "ptype": get("Product Type"), "form": get("Form Factor"),
            "qty": get("Antenna Quantity"), "mount": get("Mounting Type"),
            "dim": get("Dimensions(mm)"), "atype": get("Antenna Type"),
            "freq": freq_text, "ranges": parse_ranges(freq_text),
            "techs": techs, "lna": get("LNA Gain(dB)"), "cables": cables,
            "conn": [c for c in (get("Connector Type1"), get("Connector Type2")) if c],
            "evb": get("SMD EVB OC"), "evbDim": get("EVB Dimensions(mm)"),
            "ip": get("IP Rating"), "ik": get("IK Rating"), "flame": get("Flame Rating"),
            "uv": get("UV Resistant"), "env": get("Environmental"),
            "temp": get("Operation Temperature"),
            "jp": (get("Compatible with Japan Market") or "").upper() == "YES",
            "replace": get("To Replace"),
            "base": get("Based Standard OC"),
        }
        if rnum in imgs and oc not in images:
            t = to_thumb(imgs[rnum])
            if t:
                images[oc] = t
        products.append({k: v for k, v in p.items() if v not in (None, [], "")})
    return products


def read_pro_jp_modules(ws) -> dict[str, set[str]]:
    """PRO-JP：回傳 {OC: {相容模組}}，模組欄位為 'Compatible with Japan Market' 之後的欄。"""
    rows = list(ws.iter_rows(values_only=True))
    hdr = list(rows[0])
    oc_i = hdr.index("OC")
    start = hdr.index("Compatible with Japan Market") + 1
    mods = {i: norm_module(hdr[i]) for i in range(start, len(hdr)) if clean(hdr[i])}
    out = defaultdict(set)
    for r in rows[2:]:
        oc = clean(r[oc_i])
        if not oc:
            continue
        for i, m in mods.items():
            if i < len(r) and clean(r[i]) == "√":
                out[oc].add(m)
    return out


# --------------------------------------------------------------------------- #
# 日本市場表
# --------------------------------------------------------------------------- #
JP_SHEETS = {"5G": "5G", "LTE (4G)": "LTE", "Wi-Fi & BT": "Wi-Fi/BT"}
JP_FIXED = ["External/Patch/Embedded", "Quectel OC", "Frequency Band (MHz)", "Technology",
            "Form Factor", "Cable Length", "IP Rating", "Connector Type", "Mounting Type",
            "Dimension (mm)", "Datasheet", "Datasheet Link", "Image", "Note"]


def parse_status(note: str | None):
    """從 Note 欄判斷 NRND / EOL 與建議替代料號。"""
    if not note:
        return None, None
    status = "EOL" if "EOL" in note.upper() else "NRND" if "NRND" in note.upper() else None
    m = re.search(r"(?:suggested|replaced with)\s+([A-Z0-9/\-]+)", note, re.I)
    return status, (m.group(1) if m else None)


def read_jp(wb, imgs_by_sheet, images):
    products: dict[str, dict] = {}
    for sheet, cat in JP_SHEETS.items():
        rows = list(wb[sheet].iter_rows(values_only=True))
        hdr = [clean(h) for h in rows[0]]
        idx = {h: i for i, h in enumerate(hdr) if h}
        mod_cols = {i: norm_module(h) for i, h in enumerate(hdr) if h and h not in JP_FIXED}
        imgs = imgs_by_sheet.get(sheet, {})
        for rnum, r in enumerate(rows[1:], start=1):
            get = lambda k: clean(r[idx[k]]) if k in idx and idx[k] < len(r) else None
            oc = get("Quectel OC")
            if not oc:
                continue
            note = get("Note")
            status, suggest = parse_status(note)
            freq = get("Frequency Band (MHz)")
            # 未標單位的 JP 頻段視為 MHz
            p = products.setdefault(oc, {"oc": oc, "cats": [], "modules": set()})
            if cat not in p["cats"]:
                p["cats"].append(cat)
            p.update({k: v for k, v in {
                "kind": get("External/Patch/Embedded"), "freq": freq,
                "ranges": parse_ranges(freq), "tech": get("Technology"),
                "form": get("Form Factor"), "cable": get("Cable Length"),
                "ip": get("IP Rating"), "conn": get("Connector Type"),
                "mount": get("Mounting Type"), "dim": get("Dimension (mm)"),
                "note": note, "status": status, "suggest": suggest,
            }.items() if v is not None and k not in p})
            for i, m in mod_cols.items():
                if i < len(r) and clean(r[i]) == "√":
                    p["modules"].add(m)
            if rnum in imgs and oc not in images:
                t = to_thumb(imgs[rnum])
                if t:
                    images[oc] = t
    return products


# --------------------------------------------------------------------------- #
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--global-xlsx", required=True, type=Path)
    ap.add_argument("--jp-xlsx", required=True, type=Path)
    ap.add_argument("--template", default=Path("template.html"), type=Path)
    ap.add_argument("--out", default=Path("index.html"), type=Path)
    ap.add_argument("--version", required=True, help="App 版本號，例如 2.0.0")
    ap.add_argument("--sw", default=Path("sw.js"), type=Path, help="要同步更新 CACHE_VERSION 的 sw.js")
    a = ap.parse_args()
    if not re.fullmatch(r"\d+\.\d+\.\d+", a.version):
        sys.exit("--version 格式需為 X.Y.Z，例如 2.0.0")
    built = dt.datetime.now().astimezone().strftime("%Y-%m-%d %H:%M %z")

    images: dict[str, str] = {}

    # ---- Global ----
    gwb = load_workbook(a.global_xlsx, read_only=True, data_only=True)
    gimg = SheetImages(a.global_xlsx)
    glob = []
    for sheet, src in (("PRO", "PRO"), ("Non-PRO", "Non-PRO")):
        glob += read_selector_sheet(gwb[sheet], gimg.row_images(sheet), src, images)
    by_oc = {p["oc"]: p for p in glob}
    pro_jp_mods = read_pro_jp_modules(gwb["PRO-JP"])

    # ---- Japan ----
    jwb = load_workbook(a.jp_xlsx, read_only=True, data_only=True)
    jimg = SheetImages(a.jp_xlsx)
    jp = read_jp(jwb, {s: jimg.row_images(s) for s in JP_SHEETS}, images)

    # 合併 PRO-JP 模組相容性 (聯集)；補入 JP 檔沒有、但 PRO-JP 有勾選的料號
    tech_to_cat = {"5G": "5G", "4G": "LTE", "WIFI": "Wi-Fi/BT"}
    for oc, mods in pro_jp_mods.items():
        if not mods:
            continue
        if oc in jp:
            jp[oc]["modules"] |= mods
            continue
        g = by_oc.get(oc)
        if not g:
            continue
        jp[oc] = {
            "oc": oc, "cats": [tech_to_cat.get(g.get("type"), g.get("type") or "Other")],
            "modules": set(mods), "kind": "External" if g.get("ptype", "").startswith("External") else "Embedded",
            "freq": g.get("freq"), "ranges": g.get("ranges", []), "tech": g.get("type"),
            "form": g.get("form"), "ip": g.get("ip"), "conn": ", ".join(g.get("conn", [])) or None,
            "mount": g.get("mount"), "dim": g.get("dim"),
            "cable": ", ".join(filter(None, (c.get("len") for c in g.get("cables", [])))) or None,
            "fromSelector": True,
        }
    # 以 Global 規格補充日本料號 (產品名稱、效率、增益)
    jp_list = []
    for p in jp.values():
        g = by_oc.get(p["oc"])
        if g:
            p["name"] = g.get("name")
            p["techs"] = g.get("techs")
            p["atype"] = g.get("atype")   # Antenna Type (Monopole / Dipole / IFA ...)
        p["modules"] = sorted(p["modules"])
        jp_list.append({k: v for k, v in p.items() if v not in (None, [], "")})

    modules = sorted({m for p in jp_list for m in p.get("modules", [])})
    data = {
        "version": a.version, "built": built,
        "builtFrom": {"global": a.global_xlsx.name, "jp": a.jp_xlsx.name},
        "global": glob, "japan": jp_list, "jpModules": modules, "images": images,
    }
    payload = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
    tpl = a.template.read_text(encoding="utf8")

    # 內嵌字型：fonts/ibm-plex-sans-latin-<weight>-normal.woff2 → @font-face data URI
    # (App 不依賴任何外部網域，離線時字型仍一致)
    font_css = []
    font_dir = a.template.parent / "fonts"
    for f in sorted(font_dir.glob("ibm-plex-sans-latin-*-normal.woff2")):
        weight = re.search(r"-(\d{3})-normal", f.name).group(1)
        b64 = base64.b64encode(f.read_bytes()).decode()
        font_css.append(
            "@font-face{font-family:'IBM Plex Sans';font-style:normal;font-display:swap;"
            f"font-weight:{weight};src:url(data:font/woff2;base64,{b64}) format('woff2');"
            "unicode-range:U+0000-00FF,U+0131,U+0152-0153,U+02BB-02BC,U+02C6,U+02DA,U+02DC,"
            "U+2000-206F,U+2074,U+20AC,U+2122,U+2191,U+2193,U+2212,U+2215,U+FEFF,U+FFFD;}")
    if not font_css:
        print("警告：找不到 fonts/*.woff2，將使用系統字型", file=sys.stderr)
    tpl = tpl.replace("/*__FONTS__*/", "\n".join(font_css))
    if "/*__DATA__*/null" not in tpl:
        sys.exit("template.html 缺少 /*__DATA__*/null 佔位符")
    a.out.write_text(tpl.replace("/*__DATA__*/null", payload.replace("</", "<\\/")), encoding="utf8")

    # version.json：App 以 no-store 讀取，用來比對是否有新版本
    a.out.with_name("version.json").write_text(
        json.dumps({"version": a.version, "built": built}, ensure_ascii=False), encoding="utf8")

    # sw.js：同步 CACHE_VERSION，讓已安裝的使用者取得新版
    if a.sw.exists():
        sw = a.sw.read_text(encoding="utf8")
        sw2, n = re.subn(r"const CACHE_VERSION = '[^']*';",
                         f"const CACHE_VERSION = 'antenna-selector-v{a.version}';", sw)
        if n != 1:
            sys.exit("sw.js 找不到 CACHE_VERSION 宣告")
        a.sw.write_text(sw2, encoding="utf8")

    print(f"版本: v{a.version}  建置時間: {built}")
    print(f"Global 產品: {len(glob)}  日本產品: {len(jp_list)}  JP 模組: {len(modules)}  圖片: {len(images)}")
    print(f"輸出: {a.out} ({a.out.stat().st_size/1024:.0f} KB)")


if __name__ == "__main__":
    main()
