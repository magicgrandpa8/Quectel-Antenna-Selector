#!/usr/bin/env python3
"""
apply_datasheet_links.py — 套用「每份 datasheet 都有任何人可檢視連結」的驗證結果

輸入：Windows 程式 create_and_verify_all_links.ps1 產生的 datasheet_links_all.csv
      (欄位：Datasheet_File, OneDrive_URL, Anonymous_Check, Note)
輸出：
  1. Excel 對照表 (每份 datasheet：料號、類型、檔名、網址、未登入驗證結果)
  2. 更新 data/datasheet_links.csv：只寫入「未登入可開啟 (OK)」的連結；
     NG 的料號保留原本連結並列出，請使用者處理

料號對應：以 data/datasheet_links.csv 既有的 Datasheet_File 對應 PN (一份檔案可對應多個料號，
例如 YEGB000Q1A 與 YEGN000Q1A 共用)；對不到的檔案，再以檔名中的料號比對工具資料。

用法：python3 tools/apply_datasheet_links.py --csv datasheet_links_all.csv --xlsx-out <Excel 路徑>
"""
from __future__ import annotations

import argparse
import csv
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
LINKS = ROOT / "data" / "datasheet_links.csv"


def tool_ocs() -> set[str]:
    s = (ROOT / "index.html").read_text(encoding="utf8")
    i = s.index("const DATA = ") + 13
    return {p["oc"] for p in json.loads(s[i:s.index("\n", i)].rstrip(";"))["items"]}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--csv", required=True, type=Path)
    ap.add_argument("--xlsx-out", required=True, type=Path)
    ap.add_argument("--dry-run", action="store_true", help="只產生 Excel，不更新 data/datasheet_links.csv")
    a = ap.parse_args()

    raw = LINKS.read_bytes()
    bom, crlf = raw.startswith(b"\xef\xbb\xbf"), b"\r\n" in raw
    rows = list(csv.reader(open(LINKS, encoding="utf-8-sig", newline="")))
    hdr, body = rows[0], rows[1:]
    by_file: dict[str, list[int]] = {}
    for i, r in enumerate(body):
        if r[2].strip():
            by_file.setdefault(r[2].strip(), []).append(i)
    ocs = tool_ocs()

    res = list(csv.DictReader(open(a.csv, encoding="utf-8-sig", newline="")))
    out_rows, updated, ng, unmatched = [], [], [], []
    covered = set()
    for r in res:
        f, url = r["Datasheet_File"].strip(), (r.get("OneDrive_URL") or "").strip()
        ok = (r.get("Anonymous_Check") or "").strip().upper() == "OK" and url.startswith("https://")
        idxs = by_file.get(f, [])
        if not idxs:                                   # 以檔名中的料號比對
            pns = [oc for oc in ocs if re.search(rf"(?:^|[_&]){re.escape(oc)}(?:[_&])", f)]
            for pn in pns:
                body.append([pn, "", f, ""])
                idxs.append(len(body) - 1)
            if not pns:
                unmatched.append(f)
        for i in idxs:
            pn = body[i][0]
            covered.add(pn)
            out_rows.append([pn, body[i][1], f, url, "OK" if ok else "NG", (r.get("Note") or "").strip()])
            if ok:
                if body[i][3].strip() != url:
                    body[i][3] = url
                    updated.append(pn)
            else:
                ng.append(pn)
    no_ds = sorted(oc for oc in ocs if oc not in covered)

    if not a.dry_run:
        with open(LINKS, "w", encoding="utf-8-sig" if bom else "utf-8", newline="") as fo:
            csv.writer(fo, lineterminator="\r\n" if crlf else "\n").writerows([hdr] + body)

    # Excel 對照表
    wb = Workbook()
    ws = wb.active
    ws.title = "Datasheet Links"
    head = ["PN", "Type", "Datasheet_File", "URL (anyone can view)", "Anonymous Check", "Note"]
    ws.append(head)
    for r in sorted(out_rows, key=lambda x: x[0]):
        ws.append(r)
    for pn in no_ds:
        ws.append([pn, "", "", "", "—", "No datasheet in the customer share folder"])
    bold, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F1F1F")
    for c in ws[1]:
        c.font, c.fill = bold, fill
    red = PatternFill("solid", fgColor="FDE2E1")
    for row in ws.iter_rows(min_row=2):
        if row[4].value == "NG":
            for c in row:
                c.fill = red
        if row[3].value:
            row[3].hyperlink = row[3].value
            row[3].font = Font(color="0563C1", underline="single")
    for i, w in enumerate([16, 8, 62, 70, 16, 40], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top")
    a.xlsx_out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(a.xlsx_out)

    print(f"CSV：{len(res)} 份 datasheet｜OK {sum(1 for r in res if (r.get('Anonymous_Check') or '').upper() == 'OK')}｜NG {sum(1 for r in res if (r.get('Anonymous_Check') or '').upper() != 'OK')}")
    print(f"更新連結：{len(updated)} 個料號" + ("（dry-run，未寫入）" if a.dry_run else ""))
    print(f"未登入無法開啟 (保留原連結，需處理)：{sorted(set(ng)) or '無'}")
    print(f"檔名對不到料號：{unmatched or '無'}")
    print(f"工具中沒有 datasheet 的料號：{no_ds or '無'}")
    print(f"Excel：{a.xlsx_out}")


if __name__ == "__main__":
    main()
