#!/usr/bin/env python3
"""
release_prepare.py — 新版 Excel (例如 Rev26.11) 的第一階段處理

用法：
    python3 tools/release_prepare.py --xlsx <新版 Excel> --out <輸出 Excel> [--kit <Windows 程式 zip>] [--json <比對結果 json>]

做的事：
  1. 找出 Excel 中最新與上一版的 PRO-YY.MM 頁面 (例如 PRO-26.11 vs PRO-26.10)
  2. 比對新增 / 刪除的 OC，依 TELEC 認證分成兩組；檢查原有 OC 是否「無 TELEC → 有 TELEC」
  3. 交叉檢查：各模組工作表列出、但 PRO 認證矩陣沒有標記的 PRO 天線
  4. 寫入 Histories (G～K 欄)：直接修改工作表 XML，不經 openpyxl 存檔 (避免遺失產品照片與格式)
  5. 產生 Windows 一鍵程式 (複製新 datasheet 到客戶分享資料夾、建立「任何人可檢視」連結、
     收集新 PDF 方便上傳給 Claude)

Histories 欄位 (表頭不存在時自動新增，並為所有資料列補上框線)：
  Added OC (TELEC) | Added OC (No TELEC) | Removed OC (TELEC) | Removed OC (No TELEC) | Existing OC: No TELEC → TELEC
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import warnings
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import load_workbook

warnings.filterwarnings("ignore")

HEADERS = ["Added OC (TELEC)", "Added OC (No TELEC)", "Removed OC (TELEC)", "Removed OC (No TELEC)",
           "Existing OC: No TELEC → TELEC"]
OC_RE = re.compile(r"^Y[A-Z0-9][A-Z0-9-]{4,}$")
CERT_COL_START = 60        # PRO 頁面：第 61 欄 (index 60) 起為各模組的 TELEC 認證矩陣
SKIP_SHEETS = {"Histories", "ModuleList"}

# 本機 OneDrive 路徑 (Windows 程式使用)
ONEDRIVE_ROOT = r"C:\Users\gogol\OneDrive - 金碗國際股份有限公司"
SOURCE_REL = r"Antenna service\01 Antenna Datasheet"
SHARE_REL = r"Customer\Antenna FAE\0.For Customer\Quectel_Antenna_Selector\V3.7_20260916_Rev26.Q3"


# --------------------------------------------------------------------------- #
# 檔名：Quectel Antenna Product and TELEC Certification_V3.71_20261029_Rev26.10.xlsx
# --------------------------------------------------------------------------- #
NAME_PREFIX = "Quectel Antenna Product and TELEC Certification"


def standard_name(src_name: str, rev: str) -> str:
    """依固定格式命名輸出 Excel。上傳時系統常把空白與點換成底線 (例如 V3_7_1、Rev26_10)，
    這裡還原為 V3.71 (主版號.次版號+修訂號)、日期 8 碼、版次取自 PRO 頁面 (Rev26.11)。"""
    m = re.search(r"V(\d+)(?:[._](\d+))?(?:[._](\d+))?(?=[_ ]\d{8})", src_name) or re.search(r"V(\d+)(?:[._](\d+))?(?:[._](\d+))?", src_name)
    ver = (m.group(1) + ("." + (m.group(2) or "") + (m.group(3) or "") if m.group(2) else "")) if m else ""
    d = re.search(r"(?<!\d)(20\d{6})(?!\d)", src_name)
    parts = [NAME_PREFIX] + ([f"V{ver}"] if ver else []) + ([d.group(1)] if d else []) + [rev]
    return "_".join(parts) + ".xlsx"


# --------------------------------------------------------------------------- #
# 1～3. 讀取與比對
# --------------------------------------------------------------------------- #
def pro_sheets(names: list[str]) -> list[str]:
    pro = [n for n in names if re.fullmatch(r"PRO-\d{2}\.\d{2}", n)]
    return sorted(pro, key=lambda n: tuple(int(x) for x in n[4:].split(".")))


def read_pro(wb, sheet: str) -> dict[str, dict]:
    rows = list(wb[sheet].iter_rows(values_only=True))
    hdr = list(rows[0])
    oci, ti = hdr.index("OC"), hdr.index("Type")
    mods = [(i, str(h).strip()) for i, h in enumerate(hdr) if i >= CERT_COL_START and h]
    out = {}
    for r in rows[2:]:                      # 第 2 列為圖例
        oc = str(r[oci]).strip() if r[oci] else ""
        if oc:
            out[oc] = {"type": r[ti], "certs": sorted({m for i, m in mods if i < len(r) and r[i] not in (None, "")})}
    return out


def module_sheet_check(wb, latest: dict[str, dict], pro_names: list[str]) -> list[dict]:
    """模組工作表列出、但 PRO 矩陣沒有標記的「PRO 天線」(真正需要注意的不一致)。"""
    issues = []
    skip = SKIP_SHEETS | set(pro_names)
    for ws in wb.worksheets:
        if ws.title in skip:
            continue
        mod = ws.title.replace("&amp;", "&")
        listed = {str(v).strip() for r in ws.iter_rows(values_only=True) for v in r if v and OC_RE.match(str(v).strip())}
        missing = sorted(oc for oc in listed if oc in latest and mod not in latest[oc]["certs"])
        if missing:
            issues.append({"module": mod, "pro_oc_not_marked_in_matrix": missing})
    return issues


def compare(xlsx: Path) -> dict:
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    pro = pro_sheets(wb.sheetnames)
    if len(pro) < 2:
        sys.exit(f"需要至少兩個 PRO-YY.MM 頁面才能比對，目前只有：{pro}")
    prev_name, new_name = pro[-2], pro[-1]
    prev, new = read_pro(wb, prev_name), read_pro(wb, new_name)
    added, removed = sorted(set(new) - set(prev)), sorted(set(prev) - set(new))
    common = set(new) & set(prev)
    res = {
        "new_sheet": new_name, "prev_sheet": prev_name,
        "new_count": len(new), "prev_count": len(prev),
        "added_telec": [oc for oc in added if new[oc]["certs"]],
        "added_no_telec": [oc for oc in added if not new[oc]["certs"]],
        "removed_telec": [oc for oc in removed if prev[oc]["certs"]],
        "removed_no_telec": [oc for oc in removed if not prev[oc]["certs"]],
        "became_telec": sorted(oc for oc in common if not prev[oc]["certs"] and new[oc]["certs"]),
        "lost_telec": sorted(oc for oc in common if prev[oc]["certs"] and not new[oc]["certs"]),
        "new_module_certs": {oc: sorted(set(new[oc]["certs"]) - set(prev[oc]["certs"]))
                             for oc in sorted(common) if set(new[oc]["certs"]) - set(prev[oc]["certs"])},
        "removed_module_certs": {oc: sorted(set(prev[oc]["certs"]) - set(new[oc]["certs"]))
                                 for oc in sorted(common) if set(prev[oc]["certs"]) - set(new[oc]["certs"])},
        "added_types": {oc: new[oc]["type"] for oc in added},
        "telec_antennas": {"prev": sum(1 for v in prev.values() if v["certs"]), "new": sum(1 for v in new.values() if v["certs"])},
        "module_sheet_issues": module_sheet_check(wb, new, pro),
    }
    res["added"] = res["added_telec"] + res["added_no_telec"]
    return res


# --------------------------------------------------------------------------- #
# 4. 寫入 Histories (XML)
# --------------------------------------------------------------------------- #
def col_num(col: str) -> int:
    n = 0
    for ch in col:
        n = n * 26 + ord(ch) - 64
    return n


def col_letter(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def sheet_xml_path(z: zipfile.ZipFile, name: str) -> str:
    wbx = z.read("xl/workbook.xml").decode("utf8")
    rels = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', z.read("xl/_rels/workbook.xml.rels").decode("utf8")))
    for n, rid in re.findall(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wbx):
        if n == name:
            t = rels[rid]
            return t.lstrip("/") if t.startswith("/") else "xl/" + t
    raise KeyError(name)


def update_histories(src: Path, out: Path, diff: dict) -> dict:
    # 以 openpyxl (唯讀) 找出表頭列、目標列 (最新版次) 與既有欄位
    wb = load_workbook(src, data_only=True)        # 一般模式：每個儲存格有確實的列號 (只讀取，不存檔)
    ws = wb["Histories"]
    row_no = []                                     # rows[i] 對應的實際列號
    rows = []
    for r in ws.iter_rows():
        row_no.append(r[0].row)
        rows.append(tuple(c.value for c in r))
    hdr_i = next(i for i, r in enumerate(rows) if r and any(str(v).strip() == "Version" for v in r if v))
    hdr = [str(v).strip() if v is not None else "" for v in rows[hdr_i]]
    vcol = hdr.index("Version")
    dcol = hdr.index("Updated Date") if "Updated Date" in hdr else None
    data = [(i, r) for i, r in enumerate(rows[hdr_i + 1:], start=hdr_i + 1) if vcol < len(r) and r[vcol]]
    want = "Rev" + diff["new_sheet"][4:]                       # PRO-26.11 → Rev26.11
    tgt = next((i for i, r in data if str(r[vcol]).strip() == want), None)
    if tgt is None:                                            # 找不到對應版次：取日期最新 (相同時取最後) 的一列
        tgt = max(data, key=lambda x: (str(x[1][dcol]) if dcol is not None and dcol < len(x[1]) else "", x[0]))[0]
    hdr_row, tgt_row = row_no[hdr_i], row_no[tgt]
    tgt_version = str(rows[tgt][vcol]).strip()
    last_row = row_no[-1]

    z = zipfile.ZipFile(src)
    sp = sheet_xml_path(z, "Histories")
    x = z.read(sp).decode("utf8")
    st = z.read("xl/styles.xml").decode("utf8")

    # 欄位：沿用既有表頭；沒有則接在最後一個表頭之後
    existing = {h: col_letter(i + 1) for i, h in enumerate(hdr) if h in HEADERS}
    next_col = max([i + 1 for i, h in enumerate(hdr) if h] + [0]) + 1
    cols = {}
    for h in HEADERS:
        if h in existing:
            cols[h] = existing[h]
        else:
            cols[h] = col_letter(next_col)
            next_col += 1
    new_headers = [h for h in HEADERS if h not in existing]

    # 樣式：參考表頭與資料列既有樣式；新增「自動換行、靠上」與 A～F「靠上」的樣式
    def cell_style(ref: str):
        m = re.search(rf'<c r="{ref}"[^>]*?\ss="(\d+)"', x)
        return int(m.group(1)) if m else 0
    first_hdr = col_letter(hdr.index("Version") + 1)
    head_s = cell_style(f"{first_hdr}{hdr_row}")
    body_s = cell_style(f"{first_hdr}{tgt_row}")
    m = re.search(r'<cellXfs count="(\d+)">(.*?)</cellXfs>', st, re.S)
    xfs = re.findall(r"<xf [^>]*/>|<xf [^>]*>.*?</xf>", m.group(2), re.S)
    n = int(m.group(1))
    added_xf = []

    def variant(s: int, wrap: bool) -> int:
        xf = xfs[s]
        attrs = re.match(r"<xf ([^>]*?)/?>", xf).group(1).rstrip("/").strip()
        attrs = re.sub(r'\sapplyAlignment="[^"]*"', "", attrs) + ' applyAlignment="1"'
        al = re.search(r"<alignment ([^/]*)/>", xf)
        a = re.sub(r'\s?(vertical|wrapText)="[^"]*"', "", al.group(1)).strip() if al else ""
        a = (a + ' vertical="top"' + (' wrapText="1"' if wrap else "")).strip()
        added_xf.append(f'<xf {attrs}><alignment {a}/></xf>')
        return n + len(added_xf) - 1
    wrap_s = variant(body_s, True)
    top_cache: dict[int, int] = {}

    def row_cells(rn: int):
        mm = re.search(rf'<row r="{rn}"([^>]*?)(/>|>(.*?)</row>)', x, re.S)
        if not mm:
            return None, {}
        cells = {}
        for c in re.finditer(r'<c r="([A-Z]+)\d+"[^>]*?(?:/>|>.*?</c>)', mm.group(3) or "", re.S):
            cells[c.group(1)] = c.group(0)
        return mm, cells

    def write_row(rn: int, updates: dict[str, str], ht: float | None = None, top_existing: bool = False):
        nonlocal x
        mm, cells = row_cells(rn)
        if mm is None:
            return
        if top_existing:
            for col, c in list(cells.items()):
                if col in cols.values():
                    continue
                s = int(re.search(r'\ss="(\d+)"', c).group(1)) if re.search(r'\ss="(\d+)"', c) else 0
                if s not in top_cache:
                    top_cache[s] = variant(s, False)
                cells[col] = re.sub(r'\ss="\d+"', f' s="{top_cache[s]}"', c) if ' s="' in c else c.replace(f'<c r="{col}{rn}"', f'<c r="{col}{rn}" s="{top_cache[s]}"', 1)
        cells.update(updates)
        attrs = mm.group(1)
        last = max(col_num(c) for c in cells)
        attrs = re.sub(r'\sspans="[^"]*"', "", attrs) + f' spans="1:{last}"'
        if ht:
            attrs = re.sub(r'\s(ht|customHeight)="[^"]*"', "", attrs) + f' ht="{ht}" customHeight="1"'
        body = "".join(cells[c] for c in sorted(cells, key=col_num))
        x = x.replace(mm.group(0), f'<row r="{rn}"{attrs}>{body}</row>', 1)

    def istr(ref: str, text: str, s: int) -> str:
        return f'<c r="{ref}" s="{s}" t="inlineStr"><is><t xml:space="preserve">{escape(text)}</t></is></c>'

    # 表頭 (只新增缺少的) 與所有資料列的框線
    if new_headers:
        write_row(hdr_row, {cols[h]: istr(f"{cols[h]}{hdr_row}", h, head_s) for h in new_headers})
        for rn in range(hdr_row + 1, last_row + 1):
            if rn != tgt_row:
                _, cells = row_cells(rn)
                if cells:
                    write_row(rn, {cols[h]: f'<c r="{cols[h]}{rn}" s="{body_s}"/>' for h in new_headers if cols[h] not in cells})
    # 目標列的值
    vals = {
        "Added OC (TELEC)": diff["added_telec"], "Added OC (No TELEC)": diff["added_no_telec"],
        "Removed OC (TELEC)": diff["removed_telec"], "Removed OC (No TELEC)": diff["removed_no_telec"],
        "Existing OC: No TELEC → TELEC": diff["became_telec"],
    }
    lines = max(len(v) for v in vals.values()) or 1
    write_row(tgt_row, {cols[h]: istr(f"{cols[h]}{tgt_row}", "\n".join(v) if v else "—", wrap_s) for h, v in vals.items()},
              ht=max(15.0, 13.5 * lines) if lines > 1 else None, top_existing=True)

    # 欄寬與範圍
    for h in new_headers:
        c = col_num(cols[h])
        if re.search(rf'<col min="{c}" max="{c}"', x):
            x = re.sub(rf'<col min="{c}" max="{c}" width="[^"]*"([^>]*)/>', rf'<col min="{c}" max="{c}" width="26.77734375"\1/>', x, count=1)
        else:
            x = x.replace("</cols>", f'<col min="{c}" max="{c}" width="26.77734375" customWidth="1"/></cols>', 1)
    last_col = max(col_num(c) for c in cols.values())
    x = re.sub(r'<dimension ref="([A-Z]+)(\d+):([A-Z]+)(\d+)"/>',
               lambda d: f'<dimension ref="{d.group(1)}{d.group(2)}:{col_letter(max(col_num(d.group(3)), last_col))}{d.group(4)}"/>', x, count=1)
    st = st.replace(m.group(0), f'<cellXfs count="{n + len(added_xf)}">' + m.group(2) + "".join(added_xf) + "</cellXfs>")

    # 重新封裝：保持原檔案順序；每個項目建立新的 ZipInfo (沿用原項目資訊會造成 CRC 錯誤)
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for info in z.infolist():
            data = x.encode("utf8") if info.filename == sp else st.encode("utf8") if info.filename == "xl/styles.xml" else z.read(info.filename)
            ni = zipfile.ZipInfo(info.filename, date_time=info.date_time)
            ni.compress_type = zipfile.ZIP_DEFLATED
            ni.external_attr = info.external_attr
            zo.writestr(ni, data)
    # 驗證：只有 Histories 與 styles.xml 改變、檔案完整
    zo = zipfile.ZipFile(out)
    changed = [nm for nm in z.namelist() if z.read(nm) != zo.read(nm)]
    assert zo.testzip() is None and set(changed) <= {sp, "xl/styles.xml"}, changed
    return {"target_row": tgt_row, "target_version": tgt_version, "columns": cols, "new_headers": new_headers,
            "media": sum(1 for nm in zo.namelist() if nm.startswith("xl/media/"))}


# --------------------------------------------------------------------------- #
# 5. Windows 一鍵程式
# --------------------------------------------------------------------------- #
PS_TEMPLATE = r'''# ============================================================================
#  {REV} 新增天線 datasheet：複製到客戶分享資料夾 + 建立「任何人可檢視」連結 + 收集 PDF
#  1) 在「01 Antenna Datasheet」(含子資料夾) 依料號尋找 datasheet：
#     英文版優先 (Datasheet)、正式版優先 (排除 Preliminary)、版本號最新
#  2) 複製到客戶分享資料夾 (已存在則略過，不覆蓋)
#  3) 以你的 Microsoft 365 帳號登入，為每份檔案建立「任何人可檢視」連結，並以未登入身分驗證可開啟
#  4) 結果：桌面「{REV}_for_Claude」資料夾 = 新 PDF + datasheet_links CSV，全部上傳給 Claude 即可
# ============================================================================
$ErrorActionPreference = 'Stop'
$Root        = '{ROOT}'
$Source      = Join-Path $Root '{SOURCE}'
$Destination = Join-Path $Root '{SHARE}'
$GraphFolder = '{GRAPH_FOLDER}'
$UploadDir   = Join-Path ([Environment]::GetFolderPath('Desktop')) '{REV}_for_Claude'
$Items = @(
{ITEMS}
)

function Find-Datasheet([string]$oc) {{
    $pat = '(?:^|[_&])' + [regex]::Escape($oc) + '(?:[_&])'
    $c = Get-ChildItem -LiteralPath $Source -Recurse -File -Filter '*.pdf' -ErrorAction SilentlyContinue | Where-Object {{ $_.Name -match $pat }}
    if (-not $c) {{ return $null }}
    $c | Sort-Object @{{ Expression = {{ if ($_.Name -match '^Preliminary') {{ 1 }} else {{ 0 }} }} }},
                     @{{ Expression = {{ if ($_.Name -match 'Datasheet') {{ 0 }} else {{ 1 }} }} }},
                     @{{ Expression = {{ if ($_.Name -match '_V(\d+)[._](\d+)') {{ - ([int]$Matches[1] * 1000 + [int]$Matches[2]) }} else {{ 0 }} }} }} |
         Select-Object -First 1
}}
function Encode-Path([string]$p) {{ ($p -split '/' | ForEach-Object {{ [uri]::EscapeDataString($_) }}) -join '/' }}
# 未登入驗證：全新工作階段 (沒有 cookie、沒有帳號) 開啟連結；最後停在登入頁 = 不是公開連結
function Test-Anonymous([string]$url) {{
    try {{
        $s = New-Object Microsoft.PowerShell.Commands.WebRequestSession
        $r = Invoke-WebRequest -Uri $url -WebSession $s -UseBasicParsing -MaximumRedirection 10 -TimeoutSec 60
        $final = if ($r.BaseResponse.ResponseUri) {{ $r.BaseResponse.ResponseUri.AbsoluteUri }} else {{ $r.BaseResponse.RequestMessage.RequestUri.AbsoluteUri }}
        if ($final -match 'login\.microsoftonline\.com|_layouts/15/Authenticate|AccessDenied') {{ return 'NG' }}
        if ($r.StatusCode -eq 200) {{ return 'OK' }} else {{ return 'NG' }}
    }} catch {{ return 'NG' }}
}}

Write-Host ''
Write-Host '=== {REV} datasheet：複製 + 建立連結 ===' -ForegroundColor Cyan
foreach ($d in @($Source, $Destination)) {{
    if (-not (Test-Path -LiteralPath $d)) {{ Write-Host "找不到資料夾 (請確認 OneDrive 已同步)：$d" -ForegroundColor Red; Read-Host '按 Enter 結束'; exit 1 }}
}}
New-Item -ItemType Directory -Force -Path $UploadDir | Out-Null

# 1-2. 尋找、複製
$found = @()
foreach ($it in $Items) {{
    $f = Find-Datasheet $it.PN
    if (-not $f) {{ Write-Host "  [找不到 datasheet] $($it.PN)" -ForegroundColor Red; $it.File = ''; continue }}
    $it.File = $f.Name
    $dst = Join-Path $Destination $f.Name
    if (Test-Path -LiteralPath $dst) {{ Write-Host "  [分享資料夾已有] $($f.Name)" -ForegroundColor Yellow }}
    else {{ Copy-Item -LiteralPath $f.FullName -Destination $dst; Write-Host "  [已複製] $($f.Name)" -ForegroundColor Green }}
    Copy-Item -LiteralPath $f.FullName -Destination (Join-Path $UploadDir $f.Name) -Force
    $found += $it
}}

# 3. 建立連結 (Microsoft Graph；第一次執行會自動安裝登入模組)
$links = @{{}}; $checks = @{{}}
try {{
    if (-not (Get-Module -ListAvailable -Name Microsoft.Graph.Authentication)) {{
        Write-Host '第一次執行：安裝 Microsoft.Graph.Authentication 模組 (約 1 分鐘)…'
        Install-Module Microsoft.Graph.Authentication -Scope CurrentUser -Force -AllowClobber
    }}
    Import-Module Microsoft.Graph.Authentication
    Write-Host '請在跳出的視窗中，用公司帳號登入並同意權限…'
    Connect-MgGraph -Scopes 'Files.ReadWrite' -NoWelcome
    Write-Host '等待 OneDrive 同步新檔案 (20 秒)…'; Start-Sleep -Seconds 20
    foreach ($it in $found) {{
        $ok = $false
        foreach ($try in 1..6) {{
            try {{
                $item = Invoke-MgGraphRequest -Method GET -Uri ("https://graph.microsoft.com/v1.0/me/drive/root:/" + (Encode-Path "$GraphFolder/$($it.File)") + '?$select=id')
                $body = @{{ type = 'view'; scope = 'anonymous' }} | ConvertTo-Json
                $r = Invoke-MgGraphRequest -Method POST -Uri "https://graph.microsoft.com/v1.0/me/drive/items/$($item.id)/createLink" -Body $body -ContentType 'application/json'
                $links[$it.PN] = $r.link.webUrl; $ok = $true; break
            }} catch {{ Start-Sleep -Seconds 10 }}
        }}
        if ($ok) {{
            $checks[$it.PN] = Test-Anonymous $links[$it.PN]
            if ($checks[$it.PN] -eq 'OK') {{ Write-Host "  [連結 + 未登入可開啟] $($it.PN)" -ForegroundColor Green }}
            else {{ Write-Host "  [連結已建立，但未登入無法開啟] $($it.PN)" -ForegroundColor Red }}
        }} else {{ Write-Host "  [連結失敗] $($it.PN) (可能尚未同步，稍後再執行一次即可)" -ForegroundColor Red }}
    }}
    Disconnect-MgGraph | Out-Null
}} catch {{
    Write-Host "建立連結失敗：$($_.Exception.Message)" -ForegroundColor Red
    Write-Host '可改在 OneDrive 網頁手動「共用 → 任何人 → 複製連結」，貼到 CSV 的 OneDrive_URL 欄。'
}}

# 4. 輸出 CSV
$csv = Join-Path $UploadDir 'datasheet_links_{REV}.csv'
$Items | ForEach-Object {{ [pscustomobject]@{{ PN = $_.PN; Type = $_.Type; Datasheet_File = $_.File; OneDrive_URL = $links[$_.PN]; Anonymous_Check = $checks[$_.PN] }} }} |
    Export-Csv -Path $csv -NoTypeInformation -Encoding UTF8
Write-Host ''
Write-Host "完成：找到 $($found.Count)/$($Items.Count) 份 datasheet，建立 $($links.Count) 條連結，未登入可開啟 $(@($checks.Values | Where-Object {{ $_ -eq 'OK' }}).Count) 條" -ForegroundColor Cyan
Write-Host "請把「$UploadDir」裡的所有檔案 (PDF + CSV) 上傳給 Claude。"
Start-Process explorer.exe $UploadDir
Read-Host '按 Enter 結束'
'''


def build_kit(diff: dict, kit: Path, rev: str) -> None:
    items = ",\n".join(f"    @{{ PN='{oc}'; Type='{diff['added_types'].get(oc) or ''}'; File='' }}" for oc in diff["added"])
    ps = PS_TEMPLATE.format(REV=rev, ROOT=ONEDRIVE_ROOT, SOURCE=SOURCE_REL, SHARE=SHARE_REL,
                            GRAPH_FOLDER=SHARE_REL.replace("\\", "/"), ITEMS=items or "")
    bat = ("@echo off\r\nREM Double-click: copy new datasheets, create sharing links, collect files for Claude\r\n"
           f'powershell -NoProfile -ExecutionPolicy Bypass -File "%~dp0datasheet_{rev}.ps1"\r\n')
    with zipfile.ZipFile(kit, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(f"datasheet_{rev}.ps1", ("\ufeff" + ps).replace("\r\n", "\n").replace("\n", "\r\n").encode("utf8"))
        z.writestr(f"Run_datasheet_{rev}.bat", bat.encode("ascii"))


# --------------------------------------------------------------------------- #
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", required=True, type=Path, help="新版 Excel")
    ap.add_argument("--out", required=True, type=Path,
                    help="輸出資料夾 (自動以標準格式命名) 或完整檔名")
    ap.add_argument("--kit", type=Path, help="Windows 一鍵程式 zip (有新增 OC 時產生)")
    ap.add_argument("--json", type=Path, help="比對結果 json")
    a = ap.parse_args()

    diff = compare(a.xlsx)
    rev = "Rev" + diff["new_sheet"][4:]
    out = a.out / standard_name(a.xlsx.name, rev) if a.out.is_dir() else a.out
    hist = update_histories(a.xlsx, out, diff)
    diff["output_xlsx"] = str(out)
    diff["histories"] = hist
    if a.json:
        a.json.write_text(json.dumps(diff, ensure_ascii=False, indent=1), encoding="utf8")
    if a.kit and diff["added"]:
        build_kit(diff, a.kit, rev)

    print(f"比對：{diff['prev_sheet']} ({diff['prev_count']} 款) → {diff['new_sheet']} ({diff['new_count']} 款)")
    print(f"  新增 {len(diff['added'])} 款：有 TELEC {diff['added_telec'] or '—'}｜無 TELEC {diff['added_no_telec'] or '—'}")
    print(f"  刪除 {len(diff['removed_telec']) + len(diff['removed_no_telec'])} 款：有 TELEC {diff['removed_telec'] or '—'}｜無 TELEC {diff['removed_no_telec'] or '—'}")
    print(f"  原有 OC 無 TELEC → 有 TELEC：{diff['became_telec'] or '—'}｜有 → 無：{diff['lost_telec'] or '—'}")
    print(f"  原有 OC 新增模組認證：{diff['new_module_certs'] or '—'}")
    print(f"  TELEC 天線數：{diff['telec_antennas']['prev']} → {diff['telec_antennas']['new']}")
    print(f"  模組工作表與 PRO 矩陣不一致 (PRO 天線)：{diff['module_sheet_issues'] or '無'}")
    print(f"輸出 Excel：{out.name}")
    print(f"Histories：寫入第 {hist['target_row']} 列 ({hist['target_version']})，欄位 {', '.join(hist['columns'].values())}"
          + (f"；新增表頭 {len(hist['new_headers'])} 欄" if hist['new_headers'] else "；沿用既有表頭") + f"；照片 {hist['media']} 張保留")
    if a.kit and diff["added"]:
        print(f"Windows 一鍵程式：{a.kit}")


if __name__ == "__main__":
    main()
